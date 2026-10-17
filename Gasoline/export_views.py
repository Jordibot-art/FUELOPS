from django.http import HttpResponse
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from datetime import datetime, timedelta
from decimal import Decimal
from django.db import models

# Report Export Libraries
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import FuelSale, Expense, Employee


# === REPORT EXPORT FUNCTIONS ===

def create_pdf_response(filename):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

def create_excel_response(filename):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# SALES REPORT
@require_http_methods(["GET"])
def export_sales_pdf(request):
    response = create_pdf_response(f"Fuel_Sales_Report_{timezone.now().strftime('%Y%m%d')}.pdf")
    doc = SimpleDocTemplate(response, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    # Title
    title_style = styles['Heading1']
    title_style.alignment = 1
    elements.append(Paragraph("⛽ FUEL SALES REPORT", title_style))
    elements.append(Paragraph(f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    # Table Data
    sales = FuelSale.objects.select_related('fuel_type').order_by('-created_at')
    data = [['Transaction ID', 'Date', 'Fuel Type', 'Liters', 'Price/L', 'Total Amount', 'Cashier', 'Status']]
    
    total_sales = Decimal(0)
    for sale in sales:
        data.append([
            sale.transaction_id,
            sale.created_at.strftime('%Y-%m-%d %H:%M'),
            sale.fuel_type.name,
            f"{sale.liters:.2f} L",
            f"₱{sale.price_per_liter:.2f}",
            f"₱{sale.total_amount:.2f}",
            sale.cashier_name,
            sale.status
        ])
        total_sales += sale.total_amount
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.orange),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.whitesmoke, colors.white])
    ]))
    elements.append(table)
    
    # Summary
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"✅ Total Sales: ₱{total_sales:.2f}", styles['Heading2']))
    
    doc.build(elements)
    return response


@require_http_methods(["GET"])
def export_sales_excel(request):
    response = create_excel_response(f"Fuel_Sales_Report_{timezone.now().strftime('%Y%m%d')}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuel Sales"
    
    # Header Style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
    center_align = Alignment(horizontal='center', vertical='center')
    
    # Headers
    headers = ['Transaction ID', 'Date', 'Fuel Type', 'Liters', 'Price/L', 'Total Amount', 'Cashier', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    # Data
    sales = FuelSale.objects.select_related('fuel_type').order_by('-created_at')
    for row, sale in enumerate(sales, 2):
        ws.cell(row=row, column=1, value=sale.transaction_id)
        ws.cell(row=row, column=2, value=sale.created_at.strftime('%Y-%m-%d %H:%M'))
        ws.cell(row=row, column=3, value=sale.fuel_type.name)
        ws.cell(row=row, column=4, value=float(sale.liters))
        ws.cell(row=row, column=5, value=float(sale.price_per_liter))
        ws.cell(row=row, column=6, value=float(sale.total_amount))
        ws.cell(row=row, column=7, value=sale.cashier_name)
        ws.cell(row=row, column=8, value=sale.status)
    
    # Auto size columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = max_length + 2
    
    wb.save(response)
    return response


# EXPENSE REPORT
@require_http_methods(["GET"])
def export_expenses_pdf(request):
    response = create_pdf_response(f"Expenses_Report_{timezone.now().strftime('%Y%m%d')}.pdf")
    doc = SimpleDocTemplate(response, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("💸 EXPENSES REPORT", styles['Heading1']))
    elements.append(Paragraph(f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    expenses = Expense.objects.select_related('category').order_by('-expense_date')
    data = [['Expense ID', 'Date', 'Category', 'Description', 'Amount', 'Approved By', 'Status']]
    
    total_expenses = Decimal(0)
    for exp in expenses:
        data.append([
            exp.expense_id,
            exp.expense_date.strftime('%Y-%m-%d'),
            exp.category.name,
            exp.description,
            f"₱{exp.amount:.2f}",
            exp.approved_by_name,
            exp.status
        ])
        total_expenses += exp.amount
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.darkred),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.whitesmoke, colors.white])
    ]))
    elements.append(table)
    
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"✅ Total Expenses: ₱{total_expenses:.2f}", styles['Heading2']))
    
    doc.build(elements)
    return response


@require_http_methods(["GET"])
def export_expenses_excel(request):
    response = create_excel_response(f"Expenses_Report_{timezone.now().strftime('%Y%m%d')}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    
    headers = ['Expense ID', 'Date', 'Category', 'Description', 'Amount', 'Approved By', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    expenses = Expense.objects.select_related('category').order_by('-expense_date')
    for row, exp in enumerate(expenses, 2):
        ws.cell(row=row, column=1, value=exp.expense_id)
        ws.cell(row=row, column=2, value=exp.expense_date.strftime('%Y-%m-%d'))
        ws.cell(row=row, column=3, value=exp.category.name)
        ws.cell(row=row, column=4, value=exp.description)
        ws.cell(row=row, column=5, value=float(exp.amount))
        ws.cell(row=row, column=6, value=exp.approved_by_name)
        ws.cell(row=row, column=7, value=exp.status)
    
    wb.save(response)
    return response


# PROFIT & LOSS REPORT
@require_http_methods(["GET"])
def export_profit_pdf(request):
    response = create_pdf_response(f"Profit_Loss_Report_{timezone.now().strftime('%Y%m%d')}.pdf")
    doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("📉 PROFIT & LOSS STATEMENT", styles['Heading1']))
    elements.append(Paragraph(f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    # Monthly data
    today = timezone.now().date()
    data = [['Month', 'Revenue', 'Expenses', 'Net Profit']]
    
    total_revenue = Decimal(0)
    total_expenses = Decimal(0)
    
    for i in range(11, -1, -1):
        month_date = today - timedelta(days=i*30)
        month_sales = FuelSale.objects.filter(
            created_at__month=month_date.month,
            created_at__year=month_date.year,
            status='Paid'
        ).aggregate(total=models.Sum('total_amount'))['total'] or 0
        month_expenses = Expense.objects.filter(
            expense_date__month=month_date.month,
            expense_date__year=month_date.year,
            status='Approved'
        ).aggregate(total=models.Sum('amount'))['total'] or 0
        
        profit = month_sales - month_expenses
        total_revenue += month_sales
        total_expenses += month_expenses
        
        data.append([
            month_date.strftime('%b %Y'),
            f"₱{month_sales:.2f}",
            f"₱{month_expenses:.2f}",
            f"₱{profit:.2f}"
        ])
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.green),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.whitesmoke, colors.white])
    ]))
    elements.append(table)
    
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"📊 TOTAL REVENUE: ₱{total_revenue:.2f}", styles['Heading2']))
    elements.append(Paragraph(f"📊 TOTAL EXPENSES: ₱{total_expenses:.2f}", styles['Heading2']))
    elements.append(Paragraph(f"✅ NET PROFIT: ₱{total_revenue - total_expenses:.2f}", styles['Heading2']))
    
    doc.build(elements)
    return response


@require_http_methods(["GET"])
def export_profit_excel(request):
    response = create_excel_response(f"Profit_Loss_Report_{timezone.now().strftime('%Y%m%d')}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Profit & Loss"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    
    headers = ['Month', 'Revenue', 'Expenses', 'Net Profit']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    today = timezone.now().date()
    for row, i in enumerate(range(11, -1, -1), 2):
        month_date = today - timedelta(days=i*30)
        month_sales = FuelSale.objects.filter(
            created_at__month=month_date.month,
            created_at__year=month_date.year,
            status='Paid'
        ).aggregate(total=models.Sum('total_amount'))['total'] or 0
        month_expenses = Expense.objects.filter(
            expense_date__month=month_date.month,
            expense_date__year=month_date.year,
            status='Approved'
        ).aggregate(total=models.Sum('amount'))['total'] or 0
        
        ws.cell(row=row, column=1, value=month_date.strftime('%b %Y'))
        ws.cell(row=row, column=2, value=float(month_sales))
        ws.cell(row=row, column=3, value=float(month_expenses))
        ws.cell(row=row, column=4, value=float(month_sales - month_expenses))
    
    wb.save(response)
    return response


# PAYROLL REPORT
@require_http_methods(["GET"])
def export_payroll_pdf(request):
    response = create_pdf_response(f"Payroll_Report_{timezone.now().strftime('%Y%m%d')}.pdf")
    doc = SimpleDocTemplate(response, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("👷 EMPLOYEE PAYROLL REPORT", styles['Heading1']))
    elements.append(Paragraph(f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    employees = Employee.objects.all().order_by('last_name', 'first_name')
    data = [['Employee ID', 'Name', 'Role', 'Basic Pay', 'Allowance', 'Deductions', 'Net Pay', 'Status']]
    
    total_payroll = Decimal(0)
    for emp in employees:
        data.append([
            emp.employee_id,
            emp.full_name,
            emp.role,
            f"₱{emp.basic_pay:.2f}",
            f"₱{emp.allowance:.2f}",
            f"-₱{emp.deductions:.2f}",
            f"₱{emp.net_pay:.2f}",
            emp.payroll_status
        ])
        total_payroll += emp.net_pay
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.blue),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),
        ('GRID', (0,0), (-1,-1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.whitesmoke, colors.white])
    ]))
    elements.append(table)
    
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"✅ TOTAL MONTHLY PAYROLL: ₱{total_payroll:.2f}", styles['Heading2']))
    
    doc.build(elements)
    return response


@require_http_methods(["GET"])
def export_payroll_excel(request):
    response = create_excel_response(f"Payroll_Report_{timezone.now().strftime('%Y%m%d')}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    
    headers = ['Employee ID', 'Name', 'Role', 'Basic Pay', 'Allowance', 'Deductions', 'Net Pay', 'Status']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    employees = Employee.objects.all().order_by('last_name', 'first_name')
    for row, emp in enumerate(employees, 2):
        ws.cell(row=row, column=1, value=emp.employee_id)
        ws.cell(row=row, column=2, value=emp.full_name)
        ws.cell(row=row, column=3, value=emp.role)
        ws.cell(row=row, column=4, value=float(emp.basic_pay))
        ws.cell(row=row, column=5, value=float(emp.allowance))
        ws.cell(row=row, column=6, value=float(emp.deductions))
        ws.cell(row=row, column=7, value=float(emp.net_pay))
        ws.cell(row=row, column=8, value=emp.payroll_status)
    
    wb.save(response)
    return response